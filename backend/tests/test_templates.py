"""GET /api/v1/{collection}/template, and the refactor that made it possible.

build_template() was extracted out of sbh_download_template so the new route and
the legacy POST /api/downloadTemplate share one implementation. The parity test
below is what guards that extraction.
"""
import io
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from sbs_server.app import downloadTemplates
from sbs_server.app.api.registry import COLLECTIONS

from conftest import STUDY_URI, TOKEN, auth_headers, existing_study, missing_study

EXCEL_BACKED = sorted(p for p, s in COLLECTIONS.items() if s.template_type)
CHASSIS_URI = 'https://synbiohub.org/user/pmeka12/Test/e_coli/1'
LOOKUP_ROW = {'results': {'bindings': [
    {'s': {'value': CHASSIS_URI}, 'id': {'value': 'e_coli'}},
]}}
XLSM_MIME = 'application/vnd.ms-excel.sheet.macroEnabled.12'


# Every SBH_* sheet build_template looks up, across all four template types.
# The real templates ship these as named tables; the stand-in has to as well or
# expand_table_in_xlsm raises and the route answers 502.
SBH_SHEETS = (
    'SBH_chassis_collections',      # strains
    'SBH_plasmids_collections',     # strains
    'SBH_media_collection',         # sample-designs
    'SBH_chemicals_collection',     # sample-designs
    'SBH_strains_collection',       # sample-designs
    'SBH_sampledesigns_collection', # study-data
)


@pytest.fixture
def blank_template(monkeypatch, fake_sbh):
    """Stand in for the GitHub fetch, recording which template was asked for.

    Carries the SBH_* lookup sheets as named tables so every template type --
    not just resources, which short-circuits before the lookups -- reaches the
    end of build_template.
    """
    fake_sbh.register('GET', '/metadata', existing_study())
    workbook = Workbook()
    workbook.active['A1'] = 'blank'
    for name in SBH_SHEETS:
        sheet = workbook.create_sheet(name)
        sheet['A1'], sheet['B1'] = 'id', 'uri'
        sheet.add_table(Table(displayName=name, ref='A1:B2'))
    buffer = io.BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()

    requested = []

    def fake_fetch(object_type_id):
        requested.append(object_type_id)
        return payload, f'{object_type_id.rsplit(".", 1)[-1]}.xlsm'

    monkeypatch.setattr(downloadTemplates, 'fetch_template_bytes', fake_fetch)
    fake_fetch.requested = requested
    fake_fetch.payload = payload
    return fake_fetch


def test_template_downloads(client, blank_template):
    response = client.get('/api/v1/resources/template', headers=auth_headers())

    assert response.status_code == 200
    assert response.mimetype == XLSM_MIME
    assert response.data


def test_media_type_matches_the_spec(client, blank_template):
    """The spec declares the .xlsm macro-enabled type, not .xlsx."""
    response = client.get('/api/v1/resources/template', headers=auth_headers())

    assert response.mimetype == XLSM_MIME
    assert 'attachment' in response.headers['Content-Disposition']


def test_the_collection_picks_the_template(client, blank_template):
    """No ?type= param -- the collection alone determines which template."""
    client.get('/api/v1/resources/template', headers=auth_headers())

    assert blank_template.requested == ['synbio.object-type.resources']


@pytest.mark.parametrize('path', EXCEL_BACKED)
def test_every_excel_collection_maps_to_its_registry_template(
    client, blank_template, fake_sbh, path
):
    """Asserts the status too: without it this passed while three of the four
    collections answered 502 on a workbook with no SBH_* sheets."""
    from conftest import FakeResponse
    fake_sbh.register('GET', '/sparql', FakeResponse(200, json_data=LOOKUP_ROW))

    response = client.get(f'/api/v1/{path}/template', headers=auth_headers())

    assert response.status_code == 200, response.get_json()
    assert response.mimetype == XLSM_MIME
    assert blank_template.requested == [COLLECTIONS[path].template_type]


def test_lookup_rows_land_in_the_workbook(client, blank_template, fake_sbh):
    """The green path through expand_table_in_xlsm, which nothing else covers.

    A strains template must come back with the chassis the study already has,
    so the sheet can cross-reference them.
    """
    from conftest import FakeResponse
    fake_sbh.register('GET', '/sparql', FakeResponse(200, json_data=LOOKUP_ROW))

    response = client.get('/api/v1/strains/template', headers=auth_headers())
    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.data))
    sheet = workbook['SBH_chassis_collections']
    assert [sheet['A2'].value, sheet['B2'].value] == ['e_coli', CHASSIS_URI]


def test_resources_template_is_returned_byte_for_byte(client, blank_template):
    """build_template short-circuits resources instead of round-tripping it.

    openpyxl load/save is lossy for .xlsm and resources has no SBH_* sheets to
    fill, so the bytes must come back untouched. The parity test below cannot
    catch a regression here -- both of its sides call build_template.
    """
    response = client.get('/api/v1/resources/template', headers=auth_headers())

    assert response.status_code == 200
    assert response.data == blank_template.payload


def test_a_populated_template_is_not_returned_byte_for_byte(client, blank_template, fake_sbh):
    """The counterpart: strains must be rewritten, or the lookups did nothing."""
    from conftest import FakeResponse
    fake_sbh.register('GET', '/sparql', FakeResponse(200, json_data=LOOKUP_ROW))

    response = client.get('/api/v1/strains/template', headers=auth_headers())

    assert response.status_code == 200
    assert response.data != blank_template.payload


def test_legacy_route_returns_the_same_bytes(client, blank_template):
    """The refactor must not change POST /api/downloadTemplate at all."""
    new = client.get('/api/v1/resources/template', headers=auth_headers())

    params = json.dumps({
        'template_type': 'synbio.object-type.resources',
        'sbh_url': 'https://api.synbiohub.org',
        'sbh_token': TOKEN,
        'collection_url': STUDY_URI,
    }).encode()
    legacy = client.post(
        '/api/downloadTemplate',
        data={'Params': (io.BytesIO(params), 'parameters.json')},
        content_type='multipart/form-data',
    )

    assert legacy.status_code == 200
    assert legacy.mimetype == new.mimetype
    assert legacy.data == new.data


def test_template_requires_credentials(client, blank_template):
    response = client.get('/api/v1/resources/template',
                          headers={'X-Study-Collection': STUDY_URI})

    assert response.status_code == 401


def test_a_failed_fetch_is_502(client, monkeypatch, fake_sbh):
    fake_sbh.register('GET', '/metadata', existing_study())

    def boom(object_type_id):
        raise RuntimeError('github unreachable')

    monkeypatch.setattr(downloadTemplates, 'fetch_template_bytes', boom)

    response = client.get('/api/v1/resources/template', headers=auth_headers())

    assert response.status_code == 502
    assert 'template' in response.get_json()['error']


def test_template_for_a_missing_study_is_404(client, fake_sbh, monkeypatch):
    """Checked before the GitHub fetch, so a mistyped study fails fast.

    Without the check, it returned 200 with an empty template.
    """
    fetched = []
    monkeypatch.setattr(downloadTemplates, 'fetch_template_bytes',
                        lambda object_type_id: fetched.append(object_type_id))
    fake_sbh.register('GET', '/metadata', missing_study())

    response = client.get('/api/v1/resources/template', headers=auth_headers())

    assert response.status_code == 404
    assert fetched == []
