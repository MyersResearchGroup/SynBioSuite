"""Build resources_synthetic.xlsx from the public Excel-to-SBOL Resources template.

Run by hand when the fixture needs regenerating (pytest never collects this):

    python tests/fixtures/build_resources_fixture.py [path/to/Resources.xlsm]

The template is resources/templates/Resources.xlsm from SynBioDex/Excel-to-SBOL
at 5910698. Every value this script writes is made up; see README.md.
"""
import argparse
import os

import openpyxl

DEFAULT_TEMPLATE = os.path.expanduser(
    '~/SPUR/Excel-to-SBOL/resources/templates/Resources.xlsm'
)
OUTPUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), 'resources_synthetic.xlsx'
)

# The sheets the fixture converts, and the lookup sheets their columns need.
DATA_SHEETS = ('chemicals', 'promoter')
KEPT_SHEETS = {
    'welcome', 'Init', 'column_definitions', 'molecule_types', 'Ontology Terms',
    *DATA_SHEETS,
}

# excel2sbol.initialise() skips 9 rows of Init, so its sheet table has its
# header on row 10 and one sheet per row after that.
FIRST_INIT_SHEET_ROW = 11
# Init's "Lib Start Row" is 4 for both data sheets: headers on 4, data from 5.
HEADER_ROW = 4
FIRST_DATA_ROW = 5
# Ontology Terms has ~2,500 roles; keep only the ones the rows below use.
ROLES_USED = ('promoter',)

DESCRIPTION = 'Made-up value for tests'
ROWS = {
    'chemicals': {
        'Chemical ID': 'fixture_chemical_1',
        'Chemical Name': 'fixture_chemical_1',
        'Chemical Description': DESCRIPTION,
    },
    'promoter': {
        'Part ID': 'fixture_promoter_1',
        'Part Name': 'fixture_promoter_1',
        'Part Description': DESCRIPTION,
        'Sequence': 'atgcatgcatgcatgc',
    },
}


def build(template_path: str, output_path: str) -> None:
    # data_only swaps every formula for the value Excel last computed. It is
    # required: the ID column and column_definitions' Namespace URL are
    # formulas, pandas reads only cached values, and openpyxl saves none -- so
    # keeping the formulas would hand excel2sbol empty cells.
    workbook = openpyxl.load_workbook(template_path, data_only=True)
    _drop_unused_sheets(workbook)
    _keep_rows(workbook['column_definitions'], lambda row: row[0] in ('Init', *DATA_SHEETS))
    _keep_rows(workbook['Ontology Terms'], lambda row: row[0] in ROLES_USED)
    for sheet_name, values in ROWS.items():
        _write_only_row(workbook[sheet_name], values)
    workbook.save(output_path)


def _drop_unused_sheets(workbook) -> None:
    """Remove the sheets and their Init rows: initialise() reads every sheet Init lists."""
    init = workbook['Init']
    for row in range(init.max_row, FIRST_INIT_SHEET_ROW - 1, -1):
        name = init.cell(row, 1).value
        if name is not None and name not in KEPT_SHEETS:
            init.delete_rows(row)
    for name in workbook.sheetnames:
        if name not in KEPT_SHEETS:
            del workbook[name]


def _keep_rows(sheet, keep) -> None:
    """Keep the header row plus the rows `keep` accepts; values only."""
    kept = [
        row for row in sheet.iter_rows(min_row=2, values_only=True) if keep(row)
    ]
    sheet.delete_rows(2, sheet.max_row)
    for row in kept:
        sheet.append(row)


def _write_only_row(sheet, values: dict) -> None:
    """Replace the template's 300 blank data rows with one made-up row."""
    for table in list(sheet.tables):
        del sheet.tables[table]
    sheet.data_validations.dataValidation = []
    sheet.delete_rows(FIRST_DATA_ROW, sheet.max_row)

    columns = {cell.value: cell.column for cell in sheet[HEADER_ROW] if cell.value}
    missing = set(values) - set(columns)
    if missing:
        raise SystemExit(f'{sheet.title} has no column(s) {sorted(missing)}; '
                         f'has the template changed?')
    for header, value in values.items():
        sheet.cell(FIRST_DATA_ROW, columns[header], value)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('template', nargs='?', default=DEFAULT_TEMPLATE)
    args = parser.parse_args()
    build(args.template, OUTPUT)
    print(f'wrote {OUTPUT} ({os.path.getsize(OUTPUT) // 1024} KB)')


if __name__ == '__main__':
    main()
