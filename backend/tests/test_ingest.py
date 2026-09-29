"""ingest.load_document: format detection, accepts enforcement, workbook guards."""
import json
import os
import sys
from unittest import mock

import excel2sbol
import pytest
import sbol2
from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

from sbs_server.app.api import ingest
from sbs_server.app.api.errors import ApiError
from sbs_server.app.api import registry
from sbs_server.app.api.registry import COLLECTIONS

TEMPLATE_DIR = os.path.expanduser('~/SPUR/Excel-to-SBOL/resources/templates')

# Made-up rows on the public Resources template; see tests/fixtures/README.md.
SYNTHETIC_WORKBOOK = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'resources_synthetic.xlsx'
)


@pytest.fixture(autouse=True)
def _isolate_excel2sbol_env():
    """excel2sbol writes SBOL_DICTIONARY (and, after a failed lookup, COUNTER)
    into os.environ and never removes them. Keep them out of the next test."""
    with mock.patch.dict(os.environ):
        yield


def make_workbook(path, sbol_version=2, sheets=('Init', 'column_definitions')):
    """A workbook shaped the way excel2sbol.initialise() reads one."""
    wb = Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(name)
        if name == 'Init':
            ws['A1'], ws['B1'] = 'SBOL Version', sbol_version
            ws['A2'], ws['B2'] = 'Library Name', 'a'
    wb.save(path)
    return path


def upload(path, filename=None):
    return FileStorage(
        stream=open(path, 'rb'),
        filename=filename or os.path.basename(path),
    )


@pytest.fixture
def uploads(tmp_path):
    d = tmp_path / 'uploads'
    d.mkdir()
    return str(d)


def load(spec_path, file_storage, uploads):
    return ingest.load_document(COLLECTIONS[spec_path], file_storage, uploads)


# ------------------------------------------------------- format classification

@pytest.mark.parametrize('filename', ['a.xlsx', 'a.xlsm', 'A.XLSM'])
def test_excel_extensions_classify_as_excel(filename):
    kind, _ = ingest._classify(filename)
    assert kind == 'excel'


@pytest.mark.parametrize('filename', ['a.xml', 'a.rdf', 'a.sbol', 'A.XML'])
def test_sbol_extensions_classify_as_sbol(filename):
    kind, _ = ingest._classify(filename)
    assert kind == 'sbol'


@pytest.mark.parametrize('filename', ['a.csv', 'a.txt', 'a.json', 'noextension'])
def test_unknown_extension_is_400(filename):
    with pytest.raises(ApiError) as excinfo:
        ingest._classify(filename)
    assert excinfo.value.status == 400


def test_missing_file_is_400(uploads):
    with pytest.raises(ApiError) as excinfo:
        load('resources', FileStorage(filename=''), uploads)
    assert excinfo.value.status == 400


# ---------------------------------------------------------- accepts enforcement

def test_excel_into_an_sbol_only_collection_is_415(tmp_path, uploads):
    """This is the registry's `accepts` field doing its job."""
    path = make_workbook(tmp_path / 'r.xlsx')

    with pytest.raises(ApiError) as excinfo:
        load('devices', upload(path), uploads)

    assert excinfo.value.status == 415
    assert 'devices' in excinfo.value.message


@pytest.mark.parametrize(
    'path', [p for p, s in COLLECTIONS.items() if s.implemented and not s.excel_backed]
)
def test_every_sbol_only_collection_rejects_excel(tmp_path, uploads, path):
    workbook = make_workbook(tmp_path / f'{path}.xlsx')

    with pytest.raises(ApiError) as excinfo:
        load(path, upload(workbook), uploads)

    assert excinfo.value.status == 415


# ------------------------------------------------------------- workbook guards

def test_workbook_missing_required_sheets_is_400(tmp_path, uploads):
    path = make_workbook(tmp_path / 'bad.xlsx', sheets=('Sheet1',))

    with pytest.raises(ApiError) as excinfo:
        load('resources', upload(path), uploads)

    assert excinfo.value.status == 400
    assert 'Init' in excinfo.value.message


def test_workbook_declaring_sbol3_is_400_not_an_empty_upload(tmp_path, uploads):
    """excel2sbol takes the workbook's version over its caller's.

    An SBOL3 workbook converts happily and then reads back empty through sbol2,
    so without this guard the request would 'succeed' having uploaded nothing.
    """
    path = make_workbook(tmp_path / 'v3.xlsx', sbol_version=3)

    with pytest.raises(ApiError) as excinfo:
        load('resources', upload(path), uploads)

    assert excinfo.value.status == 400
    assert 'SBOL Version 3' in excinfo.value.message


def test_workbook_without_a_declared_version_is_400(tmp_path, uploads):
    wb = Workbook()
    wb.remove(wb.active)
    for name in ('Init', 'column_definitions'):
        wb.create_sheet(name)
    path = tmp_path / 'noversion.xlsx'
    wb.save(path)

    with pytest.raises(ApiError) as excinfo:
        load('resources', upload(path), uploads)

    assert excinfo.value.status == 400


def test_unreadable_sbol_is_400(tmp_path, uploads):
    path = tmp_path / 'junk.xml'
    path.write_text('this is not SBOL')

    with pytest.raises(ApiError) as excinfo:
        load('devices', upload(path), uploads)

    assert excinfo.value.status == 400


def test_empty_sbol_document_is_400(tmp_path, uploads, sbol_file):
    """A document with no top-levels must not submit silently."""
    empty = tmp_path / 'empty.xml'
    empty.write_text(
        '<?xml version="1.0" ?><rdf:RDF '
        'xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#"></rdf:RDF>'
    )

    with pytest.raises(ApiError) as excinfo:
        load('devices', upload(str(empty)), uploads)

    assert excinfo.value.status == 400


# ------------------------------------------------------------------- happy path

def test_valid_sbol_loads(uploads, sbol_file):
    doc = load('devices', upload(sbol_file), uploads)

    assert len(doc) >= 1


def test_temp_files_are_cleaned_up(uploads, sbol_file):
    load('devices', upload(sbol_file), uploads)

    assert os.listdir(uploads) == []


def test_temp_files_are_cleaned_up_on_failure(tmp_path, uploads):
    path = make_workbook(tmp_path / 'bad.xlsx', sheets=('Sheet1',))

    with pytest.raises(ApiError):
        load('resources', upload(path), uploads)

    assert os.listdir(uploads) == []


@pytest.mark.skipif(
    not os.path.isdir(TEMPLATE_DIR),
    reason=f'Excel-to-SBOL templates not present at {TEMPLATE_DIR}',
)
@pytest.mark.parametrize(
    'template, collection',
    [('Resources.xlsm', 'resources'),
     ('Strains.xlsm', 'strains'),
     ('SampleDesign.xlsm', 'sample-designs')],
)
def test_real_blank_templates_pass_validation(tmp_path, template, collection):
    """The real templates must clear the guards, whatever they convert to."""
    path = os.path.join(TEMPLATE_DIR, template)
    if not os.path.exists(path):
        pytest.skip(f'{template} not in {TEMPLATE_DIR}')

    ingest._validate_workbook(path)


# --------------------------------------------------------- real conversion
# Everything above this line either rejects a workbook or stops at validation.
# These actually run excel2sbol, which nothing else in the suite does.

FILLED_WORKBOOK = os.path.expanduser(
    '~/SPUR/Align-TF-Cytom-12-plasmids_Resources.xlsm'
)


def test_conversion_uses_the_excel_homespace_not_the_collections(tmp_path, monkeypatch):
    """Workbooks convert against EXCEL_HOMESPACE whatever collection they are for.

    route.py:xdc_run hardcodes https://example.org/ for every workbook, so the
    two upload paths have to agree. resources' own spec.homespace is different,
    which is exactly the confusion this pins down.
    """
    seen = {}

    def fake_convert(source, out, homespace):
        seen['homespace'] = homespace
        return []

    monkeypatch.setattr(ingest, 'convert_to_sbol', fake_convert)
    path = make_workbook(tmp_path / 'ok.xlsx')

    ingest._convert_workbook(str(path), str(tmp_path / 'out.xml'))

    assert seen['homespace'] == registry.EXCEL_HOMESPACE
    assert seen['homespace'] != COLLECTIONS['resources'].homespace


def synthetic_identities():
    home = registry.EXCEL_HOMESPACE
    return [
        f'{home}fixture_chemical_1/1',
        f'{home}fixture_promoter_1/1',
        f'{home}fixture_promoter_1_sequence/1',
    ]


def identities(doc):
    return sorted(tl.identity for tl in doc)


def test_synthetic_workbook_converts_offline(uploads):
    """Real excel2sbol on a committed workbook, so this one runs in CI.

    Nothing is mocked and sbol2's VALIDATE is at its default, True. If the
    conversion let sbol2 validate, doc.write() would POST the document to
    validator.sbolstandard.org, _block_network would refuse, and the
    conversion would fail (H3).
    """
    doc = load('resources', upload(SYNTHETIC_WORKBOOK), uploads)

    home = registry.EXCEL_HOMESPACE
    assert identities(doc) == synthetic_identities()
    chemical = doc.getComponentDefinition(f'{home}fixture_chemical_1/1')
    promoter = doc.getComponentDefinition(f'{home}fixture_promoter_1/1')
    assert list(chemical.types) == [sbol2.BIOPAX_SMALL_MOLECULE]
    assert list(promoter.roles) == [sbol2.SO_PROMOTER]
    assert list(promoter.sequences) == [f'{home}fixture_promoter_1_sequence/1']
    assert os.listdir(uploads) == [], 'conversion left temp files behind'


def test_conversion_turns_validation_back_on(uploads):
    load('resources', upload(SYNTHETIC_WORKBOOK), uploads)

    assert sbol2.Config.getOption(sbol2.ConfigOptions.VALIDATE) is True


def test_conversion_runs_unvalidated_and_restores_validation_when_it_fails(
    tmp_path, monkeypatch,
):
    seen = []

    def failing_converter(**kwargs):
        seen.append(sbol2.Config.getOption(sbol2.ConfigOptions.VALIDATE))
        raise RuntimeError('conversion blew up')

    monkeypatch.setattr(excel2sbol, 'converter', failing_converter)

    with pytest.raises(ApiError):
        ingest._convert_workbook(SYNTHETIC_WORKBOOK, str(tmp_path / 'out.xml'))

    assert seen == [False], 'excel2sbol ran with sbol2 validation on'
    assert sbol2.Config.getOption(sbol2.ConfigOptions.VALIDATE) is True


# ------------------------------------------------- conversion stays offline
# excel2sbol fetches URLs named in workbook cells, looks displayIds and
# sequences up on the welcome sheet's Domain, and keeps per-run state in
# os.environ that outlives the request (review DEFERRED, excel2sbol). Any
# request these tests let through trips _block_network, which fails the
# conversion -- so "converts" also means "sent nothing".

NCBI_SEQUENCE_URL = 'https://www.ncbi.nlm.nih.gov/nuccore/NC_000913.3'
EXCEL2SBOL_RUN_ENV = ('COUNTER', 'SBOL_DICTIONARY')


def synthetic_variant(tmp_path, edit):
    """resources_synthetic.xlsx with `edit` applied, saved under tmp_path."""
    workbook = load_workbook(SYNTHETIC_WORKBOOK)
    edit(workbook)
    path = tmp_path / 'variant.xlsx'
    workbook.save(path)
    return str(path)


def test_a_url_in_a_sequence_cell_is_kept_without_being_fetched(tmp_path, uploads):
    """library2.link_validation GETs every URL in a Sequence cell.

    The public template's own pattern allows an NCBI nuccore link here. It is
    stored as the sequence reference, as a reachable link always was.
    """
    def put_a_url_in_the_sequence(workbook):
        workbook['promoter']['L5'] = NCBI_SEQUENCE_URL

    path = synthetic_variant(tmp_path, put_a_url_in_the_sequence)
    doc = load('resources', upload(path), uploads)

    promoter = doc.getComponentDefinition(
        f'{registry.EXCEL_HOMESPACE}fixture_promoter_1/1'
    )
    assert list(promoter.sequences) == [NCBI_SEQUENCE_URL]


def test_a_welcome_sheet_domain_is_never_contacted(tmp_path, uploads):
    """A Domain made excel2sbol search that host for every displayId."""
    def set_a_domain(workbook):
        workbook['welcome']['C16'] = 'https://sbh.example.org'

    path = synthetic_variant(tmp_path, set_a_domain)
    doc = load('resources', upload(path), uploads)

    assert identities(doc) == synthetic_identities()


def test_a_poisoned_counter_does_not_fail_the_next_conversion(uploads, monkeypatch):
    """A 401/403/404 on a checked URL set COUNTER and nothing cleared it, so
    every later conversion in that uwsgi worker skipped writing its output."""
    monkeypatch.setenv('COUNTER', 'Error found')

    doc = load('resources', upload(SYNTHETIC_WORKBOOK), uploads)

    assert identities(doc) == synthetic_identities()


def test_a_previous_uploads_domain_is_not_used(tmp_path, uploads, monkeypatch):
    """SBOL_DICTIONARY outlived the upload that set it, so a workbook with no
    welcome sheet sent its displayIds to the previous uploader's Domain."""
    monkeypatch.setenv('SBOL_DICTIONARY', json.dumps({
        'Domain': 'https://sbh.example.org',
        'Library Name': 'previous upload',
        'Master Collection': '',
    }))

    def drop_the_welcome_sheet(workbook):
        del workbook['welcome']

    path = synthetic_variant(tmp_path, drop_the_welcome_sheet)
    doc = load('resources', upload(path), uploads)

    assert identities(doc) == synthetic_identities()


def test_conversion_leaves_no_excel2sbol_state_behind(uploads):
    load('resources', upload(SYNTHETIC_WORKBOOK), uploads)

    assert not set(EXCEL2SBOL_RUN_ENV) & set(os.environ)


def test_a_failed_conversion_leaves_no_excel2sbol_state_behind(tmp_path, monkeypatch):
    def converter_that_poisons_then_fails(**kwargs):
        os.environ['COUNTER'] = 'Error found'
        raise RuntimeError('conversion blew up')

    monkeypatch.setattr(excel2sbol, 'converter', converter_that_poisons_then_fails)

    with pytest.raises(ApiError):
        ingest._convert_workbook(SYNTHETIC_WORKBOOK, str(tmp_path / 'out.xml'))

    assert 'COUNTER' not in os.environ


def test_sys_exit_in_the_converter_is_a_400(tmp_path, monkeypatch):
    """excel2sbol calls sys.exit(1) on some input. SystemExit is not an
    Exception, so it went past every handler instead of becoming a 400."""
    link_validation = excel2sbol.library2.link_validation

    def exiting_converter(**kwargs):
        sys.exit(1)

    monkeypatch.setattr(excel2sbol, 'converter', exiting_converter)

    with pytest.raises(ApiError) as excinfo:
        ingest._convert_workbook(SYNTHETIC_WORKBOOK, str(tmp_path / 'out.xml'))

    assert excinfo.value.status == 400
    assert excel2sbol.library2.link_validation is link_validation
    assert sbol2.Config.getOption(sbol2.ConfigOptions.VALIDATE) is True


@pytest.mark.skipif(
    not os.path.isdir(TEMPLATE_DIR),
    reason=f'Excel-to-SBOL templates not present at {TEMPLATE_DIR}',
)
def test_a_real_template_survives_conversion(tmp_path):
    """excel2sbol actually runs -- not mocked, not stopped at validation.

    A blank template legitimately yields no objects; what matters is that the
    conversion completes instead of raising.
    """
    path = os.path.join(TEMPLATE_DIR, 'Strains.xlsm')
    if not os.path.exists(path):
        pytest.skip(f'Strains.xlsm not in {TEMPLATE_DIR}')

    doc = ingest._convert_workbook(path, str(tmp_path / 'out.xml'))

    assert doc is not None
    assert len(doc) == 0, 'a blank template should convert to nothing'


@pytest.mark.skipif(
    not os.path.exists(FILLED_WORKBOOK),
    reason=f'filled workbook not present at {FILLED_WORKBOOK}',
)
def test_a_filled_workbook_converts_to_real_objects(uploads):
    """The flagship path end to end: .xlsm in, populated SBOL document out."""
    doc = load('resources', upload(FILLED_WORKBOOK), uploads)

    assert len(doc) > 1, 'a filled Resources workbook should yield many objects'
    identities = [tl.identity for tl in doc]
    assert all(i.startswith(registry.EXCEL_HOMESPACE) for i in identities), \
        f'unexpected homespace: {identities[:3]}'
    assert os.listdir(uploads) == [], 'conversion left temp files behind'
