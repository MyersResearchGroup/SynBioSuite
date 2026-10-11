"""POST /api/downloadTemplate, through the HTTP client.

test_templates.py covers build_template and the route's success and SSRF
paths; this is the rest of the request handling. The tests patch only
sbh_http and requests, never the service itself, so they held unedited while
the code moved out of downloadTemplates.py. Behaviour that is odd but kept on
purpose is marked "Wart, pinned".
"""
import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from sbs_server.app import downloadTemplates

from conftest import STUDY_URI, TOKEN, FakeResponse
from legacy_helpers import (
    COLLECTION_GRAPH,
    SBH_URL,
    params_part,
    post,
    raw_part,
    requested,
    sparql_queries,
    without,
)

pytestmark = pytest.mark.usefixtures('workdir')

URL = '/api/downloadTemplate'
GITHUB = 'raw.githubusercontent.com'
STRAINS_TEMPLATE_URL = (
    'https://raw.githubusercontent.com/SynBioDex/Excel-to-SBOL/master/'
    'resources/templates/Strains.xlsm'
)
XLSM_MIME = 'application/vnd.ms-excel.sheet.macroEnabled.12'
CHASSIS_URI = 'https://synbiohub.org/user/pmeka12/Test/e_coli/1'
STRAINS_SHEETS = ('SBH_chassis_collections', 'SBH_plasmids_collections')
UNEXPECTED_FAILURE = (
    'The template download failed unexpectedly. The details are in the server log.'
)

PARAMS = {
    'template_type': 'synbio.object-type.strains',
    'sbh_url': SBH_URL,
    'sbh_token': TOKEN,
    'collection_url': STUDY_URI,
}
REQUIRED = ['template_type', 'sbh_url', 'sbh_token', 'collection_url']


def strains_template() -> bytes:
    """A stand-in for the GitHub template: the two lookup sheets, as named tables."""
    workbook = Workbook()
    for name in STRAINS_SHEETS:
        sheet = workbook.create_sheet(name)
        sheet['A1'], sheet['B1'] = 'id', 'uri'
        sheet.add_table(Table(displayName=name, ref='A1:B2'))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def one_chassis():
    return FakeResponse(200, json_data={'results': {'bindings': [
        {'s': {'value': CHASSIS_URI}, 'id': {'value': 'e_coli'}},
    ]}})


def download(client, **overrides):
    return post(client, URL, Params=params_part({**PARAMS, **overrides}))


# ------------------------------------------------------------------ success

def test_a_template_comes_back_filled_from_the_study(client, legacy_sbh):
    legacy_sbh.register('GET', GITHUB, FakeResponse(200, content=strains_template()))
    legacy_sbh.register('GET', '/sparql', one_chassis())

    response = download(client)

    assert response.status_code == 200
    assert response.mimetype == XLSM_MIME
    assert 'attachment' in response.headers['Content-Disposition']
    assert 'Strains.xlsm' in response.headers['Content-Disposition']
    sheet = load_workbook(io.BytesIO(response.data))['SBH_chassis_collections']
    assert [sheet['A2'].value, sheet['B2'].value] == ['e_coli', CHASSIS_URI]


def test_the_template_is_fetched_and_the_study_queried(client, legacy_sbh):
    legacy_sbh.register('GET', GITHUB, FakeResponse(200, content=strains_template()))
    legacy_sbh.register('GET', '/sparql', one_chassis())

    download(client)

    assert requested(legacy_sbh) == [
        f'GET {STRAINS_TEMPLATE_URL}',
        f'GET {SBH_URL}/sparql',   # chassis
        f'GET {SBH_URL}/sparql',   # plasmids
    ]
    for query in sparql_queries(legacy_sbh):
        assert f'FROM <{COLLECTION_GRAPH}>' in query
        assert f'<{STUDY_URI}> sbol:member' in query


# ------------------------------------------------- rejected before any request

def assert_rejected(response, sbh, error):
    assert response.status_code == 400
    assert response.get_json() == {'error': error}
    assert sbh.calls == []


def test_no_params_part_is_400(client, legacy_sbh):
    response = post(client, URL)

    assert_rejected(response, legacy_sbh, 'No Params file part')


def test_a_params_part_without_a_filename_is_400(client, legacy_sbh):
    response = post(client, URL, Params=params_part(PARAMS, filename=''))

    assert_rejected(response, legacy_sbh, 'No selected Params file')


@pytest.mark.parametrize('missing', REQUIRED)
def test_each_required_param_is_checked(client, legacy_sbh, missing):
    response = post(client, URL, Params=params_part(without(PARAMS, missing)))

    assert_rejected(response, legacy_sbh, f'Parameter {missing} not found in request')


def test_a_null_token_is_400(client, legacy_sbh):
    response = download(client, sbh_token=None)

    assert_rejected(response, legacy_sbh, 'No SBH credentials provided')


def test_an_unusable_sbh_url_is_400(client, legacy_sbh):
    response = download(client, sbh_url='ftp://synbiohub.org')

    assert_rejected(response, legacy_sbh, 'Invalid sbh_url: it must be an http(s) URL')


def test_the_token_is_checked_before_the_url(client, legacy_sbh):
    response = download(client, sbh_token=None, sbh_url='ftp://synbiohub.org')

    assert_rejected(response, legacy_sbh, 'No SBH credentials provided')


# ----------------------------------------------- failures while building it

def test_an_unknown_template_type_is_a_500(client, legacy_sbh):
    """Wart, pinned: the caller's mistake, answered as a server error."""
    response = download(client, template_type='nope')

    assert response.status_code == 500
    assert response.get_json() == {'error': 'Unknown object type: nope'}
    assert legacy_sbh.calls == []


def test_an_unknown_template_type_has_an_error_of_its_own():
    """What lets the route keep that message while hiding unexpected ones.
    Still a ValueError, as it was."""
    with pytest.raises(downloadTemplates.UnknownTemplateError) as excinfo:
        downloadTemplates.fetch_template_bytes('nope')

    assert isinstance(excinfo.value, ValueError)
    assert str(excinfo.value) == 'Unknown object type: nope'


def test_a_template_github_will_not_serve_is_a_500(client, legacy_sbh):
    legacy_sbh.register('GET', GITHUB, FakeResponse(404))

    response = download(client)

    assert response.status_code == 500
    assert response.get_json() == {'error': '404 error'}


def test_a_failed_lookup_is_a_500(client, legacy_sbh):
    legacy_sbh.register('GET', GITHUB, FakeResponse(200, content=strains_template()))
    legacy_sbh.register('GET', '/sparql', FakeResponse(500, text='boom'))

    response = download(client)

    assert response.status_code == 500
    assert response.get_json() == {'error': 'SynBioHub sparql query failed (500): boom'}


def test_an_unexpected_failure_tells_the_caller_nothing_about_it(
    client, legacy_sbh, caplog,
):
    """Here, GitHub serving something that is not a workbook."""
    legacy_sbh.register('GET', GITHUB, FakeResponse(200, content=b'not a workbook'))

    response = download(client)

    assert response.status_code == 500
    assert response.get_json() == {'error': UNEXPECTED_FAILURE}
    assert 'File is not a zip file' in caplog.text


# ------------------------------------------------------- unhandled, so a 500

def test_params_that_are_not_json_are_unhandled(unhandled_client, legacy_sbh):
    """Wart, pinned: Flask's own error page, because nothing catches it."""
    response = post(
        unhandled_client, URL, Params=raw_part(b'{not json', 'parameters.json'),
    )

    assert response.status_code == 500
    assert response.mimetype == 'text/html'
    assert legacy_sbh.calls == []
